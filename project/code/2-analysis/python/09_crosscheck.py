"""Tahap 2 / langkah 09: cek silang angka pesan bot vs SQL independen vs Excel (Q12).

Tiga sumber untuk 12 tanggal sampel (hari kejadian, lonjakan, data terlambat, dan sampel acak tetap):
  A. pesan bot (teks yang benar-benar dibuat, angka dibaca dengan regex),
  B. SQL independen: query langsung ke CSV bersih (tanpa view) lewat DuckDB,
  C. Excel: rumus SUMIFS di crosscheck.xlsx dihitung ulang oleh LibreOffice (BUKAN Microsoft Excel).
Keluaran: results/2-analysis/{crosscheck.csv, crosscheck.xlsx}
"""
import random
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import duckdb
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from bot.report import build_message, load_kpis  # noqa: E402
from common import CLEAN_DIR, DB_PATH, RES_ANALYSIS, SEED, STATIONS  # noqa: E402

TOL_DISPLAY = 0.0005     # angka di pesan dibulatkan 1 desimal persen -> selisih maksimal 0,05 pp


def main():
    con = duckdb.connect(str(DB_PATH), read_only=True)
    days = [r[0] for r in con.execute("SELECT strftime(date,'%Y-%m-%d') FROM v_line_daily ORDER BY date").fetchall()]
    fixed = ["2026-03-30", "2026-03-31", "2026-05-04", "2026-06-02", "2026-04-23", "2026-03-23"]
    rnd = random.Random(SEED)
    sample = sorted(set(fixed + rnd.sample([d for d in days if d not in fixed], 6)))

    prod = pd.read_csv(CLEAN_DIR / "production_daily.csv")
    dl = pd.read_csv(CLEAN_DIR / "defect_log.csv").merge(pd.read_csv(CLEAN_DIR / "dim_defect.csv")[["defect_code", "defect_name"]], on="defect_code")

    # B. SQL independen dari CSV
    sqlres = {}
    ind = duckdb.connect()
    for d in sample:
        q = ind.execute(f"""
          WITH s AS (SELECT station, SUM(ok_first_pass)*1.0/SUM(input_qty) fpy, SUM(input_qty) i, SUM(ok_first_pass) o
                     FROM read_csv_auto('{CLEAN_DIR / 'production_daily.csv'}') WHERE CAST(date AS VARCHAR)='{d}' GROUP BY 1)
          SELECT (SELECT o FROM s WHERE station='Packing')*1.0/(SELECT i FROM s WHERE station='Frame Welding'),
                 (SELECT MIN(fpy) FROM s)""").fetchone()
        top = ind.execute(f"""SELECT defect_name, SUM(qty) q FROM read_csv_auto('{CLEAN_DIR / 'defect_log.csv'}') l
                              JOIN read_csv_auto('{CLEAN_DIR / 'dim_defect.csv'}') USING (defect_code)
                              WHERE CAST(date AS VARCHAR)='{d}' GROUP BY 1 ORDER BY q DESC, defect_name LIMIT 1""").fetchone()
        sqlres[d] = dict(line=q[0], worst=q[1], top_name=top[0], top_qty=int(top[1]))

    # C. Excel
    wb = openpyxl.Workbook()
    wp = wb.active; wp.title = "production"
    wp.append(list(prod.columns)); [wp.append(list(r)) for r in prod.itertuples(index=False)]
    wd = wb.create_sheet("defect")
    cols = ["date", "defect_name", "qty"]
    wd.append(cols); [wd.append(list(r)) for r in dl[cols].itertuples(index=False)]
    wc = wb.create_sheet("check")
    n_p, n_d = len(prod) + 1, len(dl) + 1
    wc.append(["date"] + STATIONS + ["line_yield", "worst_station_yield", "top_defect_name (dari SQL)", "top_defect_qty"])
    for r, d in enumerate(sample, start=2):
        wc.cell(r, 1, d)
        for j, st in enumerate(STATIONS):
            wc.cell(r, 2 + j, f'=SUMIFS(production!$G$2:$G${n_p},production!$A$2:$A${n_p},$A{r},production!$C$2:$C${n_p},"{st}")'
                              f'/SUMIFS(production!$F$2:$F${n_p},production!$A$2:$A${n_p},$A{r},production!$C$2:$C${n_p},"{st}")')
        c_ok = f'SUMIFS(production!$G$2:$G${n_p},production!$A$2:$A${n_p},$A{r},production!$C$2:$C${n_p},"Packing")'
        c_in = f'SUMIFS(production!$F$2:$F${n_p},production!$A$2:$A${n_p},$A{r},production!$C$2:$C${n_p},"Frame Welding")'
        wc.cell(r, 2 + len(STATIONS), f"={c_ok}/{c_in}")
        wc.cell(r, 3 + len(STATIONS), f"=MIN(B{r}:{openpyxl.utils.get_column_letter(1 + len(STATIONS))}{r})")
        wc.cell(r, 4 + len(STATIONS), sqlres[d]["top_name"])
        wc.cell(r, 5 + len(STATIONS), f'=SUMIFS(defect!$C$2:$C${n_d},defect!$A$2:$A${n_d},$A{r},defect!$B$2:$B${n_d},{openpyxl.utils.get_column_letter(4 + len(STATIONS))}{r})')
    tmp = Path(tempfile.mkdtemp())
    wb.save(tmp / "crosscheck.xlsx")
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(tmp / "out"), str(tmp / "crosscheck.xlsx")],
                   check=True, capture_output=True, timeout=240)
    shutil.copy(tmp / "out" / "crosscheck.xlsx", RES_ANALYSIS / "crosscheck.xlsx")
    ws = openpyxl.load_workbook(RES_ANALYSIS / "crosscheck.xlsx", data_only=True)["check"]
    L = len(STATIONS)
    xl = {ws.cell(r, 1).value: dict(line=ws.cell(r, 2 + L).value, worst=ws.cell(r, 3 + L).value, top_qty=ws.cell(r, 5 + L).value) for r in range(2, len(sample) + 2)}

    rows = []
    for d in sample:
        k = load_kpis(con, d)
        msg = build_message(k)
        m_line = float(re.search(r"Yield lini ([\d,]+)%", msg).group(1).replace(",", ".")) / 100
        m_worst = float(re.search(r"Yield terendah: .+? ([\d,]+)%", msg).group(1).replace(",", ".")) / 100
        m_top = re.search(r"Top NG hari ini: (.+?) \((\d+)\)", msg)
        s, x = sqlres[d], xl[d]
        ok_line = abs(m_line - s["line"]) <= TOL_DISPLAY and abs(s["line"] - x["line"]) < 1e-9
        ok_worst = abs(m_worst - s["worst"]) <= TOL_DISPLAY and abs(s["worst"] - x["worst"]) < 1e-9
        ok_top = m_top.group(1) == s["top_name"] and int(m_top.group(2)) == s["top_qty"] == int(x["top_qty"])
        rows.append(dict(date=d, bot_line=m_line, sql_line=round(s["line"], 6), excel_line=round(x["line"], 6), line_pass=ok_line,
                         bot_worst=m_worst, sql_worst=round(s["worst"], 6), excel_worst=round(x["worst"], 6), worst_pass=ok_worst,
                         bot_top=f"{m_top.group(1)} ({m_top.group(2)})", sql_top=f"{s['top_name']} ({s['top_qty']})", excel_top_qty=int(x["top_qty"]), top_pass=ok_top))
    out = pd.DataFrame(rows)
    out.to_csv(RES_ANALYSIS / "crosscheck.csv", index=False)
    print(out[["date", "line_pass", "worst_pass", "top_pass"]].to_string(index=False))
    print("semua lulus:", bool(out[["line_pass", "worst_pass", "top_pass"]].all().all()))


if __name__ == "__main__":
    main()
